# app/report_pptx.py
# 운영 리포트를 PowerPoint 파일로 만든다. 자료는 app/report.py 가 모은 것을 그대로 쓰고,
# 여기서는 슬라이드에 배치하는 일만 한다.
#
# python-pptx 를 쓰는 이유: 이 함수는 웹 요청 중에 실행되므로 파이썬 안에서 끝나야 한다.
# (Node 기반 도구를 쓰면 요청마다 외부 프로세스를 띄워야 한다.)

from datetime import datetime, timezone
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# ----------------------------------------------------------------------
# 색 — 화면(style.css)에서 쓰는 것과 같은 값을 쓴다.
# 리포트와 화면의 심각도 색이 다르면 같은 데이터를 두 번 배워야 한다.
# ----------------------------------------------------------------------
INK        = RGBColor(0x1A, 0x1D, 0x21)   # 본문 글자
MUTED      = RGBColor(0x62, 0x69, 0x6F)   # 보조 글자
LINE       = RGBColor(0xE3, 0xE6, 0xEA)   # 표 구분선
DARK_BG    = RGBColor(0x16, 0x18, 0x1C)   # 표지/마무리 배경
DARK_INK   = RGBColor(0xFF, 0xFF, 0xFF)
DARK_MUTED = RGBColor(0x9A, 0xA3, 0xAD)
ACCENT     = RGBColor(0x2A, 0x78, 0xD6)
SURFACE    = RGBColor(0xF7, 0xF8, 0xFA)   # 카드 배경

SEVERITY = {
    "critical": RGBColor(0xD0, 0x3B, 0x3B),
    "error":    RGBColor(0xEC, 0x83, 0x5A),
    "warning":  RGBColor(0xFA, 0xB2, 0x19),
    "info":     RGBColor(0x6B, 0x74, 0x80),
}
UP   = RGBColor(0xB9, 0x1C, 0x1C)   # 늘어남 = 나쁜 신호
DOWN = RGBColor(0x14, 0x65, 0x32)   # 줄어듦 = 좋은 신호

# 제목은 세리프, 본문은 산세리프. 둘 다 Office 기본 포함 글꼴이라 어디서나 같게 보인다.
HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)   # 16:9
MARGIN = Inches(0.7)


# ----------------------------------------------------------------------
# 작은 도우미들
# ----------------------------------------------------------------------
def _blank(prs):
    """빈 레이아웃 슬라이드. 자리표시자 없이 직접 배치한다."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill(shape, color):
    """도형 배경을 칠하고 테두리는 없앤다."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _fill_cell(cell, color):
    """표 셀 배경. 셀에는 .line 이 없으므로 도형용 _fill 을 쓸 수 없다."""
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def _bg(slide, prs, color):
    """슬라이드 전체 배경. 사각형을 깔고 맨 뒤로 보낸다."""
    from pptx.enum.shapes import MSO_SHAPE
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _fill(box, color)
    # 맨 뒤로 (다른 도형이 위에 오도록)
    slide.shapes._spTree.remove(box._element)
    slide.shapes._spTree.insert(2, box._element)


def _text(slide, left, top, width, height, text, *, size=14, bold=False,
          color=INK, font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          spacing=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    # 도형과 글자를 같은 x 에 맞추려면 안쪽 여백을 0 으로 둬야 한다.
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing is not None:
            p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def _title(slide, text, sub=None):
    """슬라이드 제목. 제목 아래 밑줄이나 색 막대는 넣지 않는다."""
    # 제목 상자 높이(0.55)와 부제 위치(1.12)가 겹치지 않도록 둔다.
    # 0.5 + 0.55 = 1.05 < 1.12
    _text(slide, MARGIN, Inches(0.5), W - MARGIN * 2, Inches(0.55),
          text, size=30, bold=True, font=HEAD_FONT)
    if sub:
        _text(slide, MARGIN, Inches(1.12), W - MARGIN * 2, Inches(0.36),
              sub, size=12, color=MUTED)


def _card(slide, left, top, width, height, fill=SURFACE):
    from pptx.enum.shapes import MSO_SHAPE
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill(box, fill)
    box.adjustments[0] = 0.06
    return box


def _delta_text(change):
    if change is None:
        return "—", MUTED
    if change > 0:
        return f"+{change}%", UP
    if change < 0:
        return f"{change}%", DOWN
    return "0%", MUTED


def _table(slide, left, top, width, rows, headers, col_widths=None, row_h=Inches(0.32)):
    """간단한 표. python-pptx 기본 표 스타일은 요란해서 직접 칠한다."""
    n_rows, n_cols = len(rows) + 1, len(headers)
    # 길이(EMU)는 정수여야 한다. 호출부에서 나눗셈으로 폭을 계산하면
    # 파이썬 / 는 float 를 돌려주고, python-pptx 는 그걸 그대로 거부한다
    # (TypeError: value must be an integral type). 여기서 한 번 정리한다.
    shape = slide.shapes.add_table(
        n_rows, n_cols, int(left), int(top), int(width), int(row_h * n_rows)
    )
    table = shape.table

    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = int(cw)

    for r in range(n_rows):
        table.rows[r].height = row_h
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _fill_cell(cell, RGBColor(0xFF, 0xFF, 0xFF) if r else SURFACE)

            value = headers[c] if r == 0 else rows[r - 1][c]
            colour = MUTED if r == 0 else INK
            # (값, 색) 튜플이면 색을 지정한 것이다.
            if isinstance(value, tuple):
                value, colour = value

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(value)
            run.font.size = Pt(10.5)
            run.font.bold = (r == 0)
            run.font.color.rgb = colour
            run.font.name = BODY_FONT
    return table


def _style_chart(chart, *, legend=False):
    """차트를 조용하게 만든다. 기본 스타일은 격자와 범례가 과하다."""
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = BODY_FONT
        chart.legend.font.color.rgb = MUTED
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.size = Pt(10)
        axis.tick_labels.font.name = BODY_FONT
        axis.tick_labels.font.color.rgb = MUTED
        axis.has_major_gridlines = False
    chart.value_axis.has_major_gridlines = True


# ----------------------------------------------------------------------
# 슬라이드
# ----------------------------------------------------------------------
def _slide_cover(prs, data):
    slide = _blank(prs)
    _bg(slide, prs, DARK_BG)

    _text(slide, MARGIN, Inches(2.4), W - MARGIN * 2, Inches(1.0),
          "운영 리포트", size=54, bold=True, color=DARK_INK, font=HEAD_FONT)
    _text(slide, MARGIN, Inches(3.5), W - MARGIN * 2, Inches(0.5),
          f"최근 {data['days']}일", size=22, color=ACCENT)
    _text(slide, MARGIN, Inches(4.3), W - MARGIN * 2, Inches(0.9),
          f"{data['start'].strftime('%Y-%m-%d %H:%M')} ~ "
          f"{data['end'].strftime('%Y-%m-%d %H:%M')} (UTC)\n"
          f"생성 {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} (UTC)",
          size=13, color=DARK_MUTED, spacing=4)
    slide.notes_slide.notes_text_frame.text = (
        "이 리포트는 이벤트 집계와 리소스 스냅샷 비교를 자동으로 모은 것입니다."
    )


def _slide_summary(prs, data, summary):
    slide = _blank(prs)
    _title(slide, "요약", f"직전 {data['days']}일과 비교")

    stats = [
        ("전체 이벤트", data["total"], data["total_change"]),
        ("알람 대상", data["alarm_total"], data["alarm_change"]),
        ("반복 알람 종류", len(data["repeated"]), None),
        ("조용해진 출처", len(data["silent_sources"]), None),
    ]
    card_w, gap = Inches(2.85), Inches(0.24)
    top = Inches(1.75)
    for i, (label, value, change) in enumerate(stats):
        left = MARGIN + i * (card_w + gap)
        _card(slide, left, top, card_w, Inches(1.45))

        # 증감률이 오른쪽 0.95in 을 쓰므로, 숫자 상자는 그만큼 좁혀야 겹치지 않는다.
        delta_w = Inches(0.95) if change is not None else Inches(0)
        _text(slide, left + Inches(0.22), top + Inches(0.18),
              card_w - Inches(0.44) - delta_w, Inches(0.7),
              str(value), size=40, bold=True)
        _text(slide, left + Inches(0.22), top + Inches(0.92), card_w - Inches(0.44), Inches(0.3),
              label, size=11, color=MUTED)
        if change is not None:
            txt, colour = _delta_text(change)
            _text(slide, left + card_w - Inches(1.05), top + Inches(0.3),
                  Inches(0.83), Inches(0.32), txt, size=14, bold=True,
                  color=colour, align=PP_ALIGN.RIGHT)

    body_top = Inches(3.5)
    if data["small_sample"]:
        _card(slide, MARGIN, body_top, W - MARGIN * 2, Inches(0.75),
              fill=RGBColor(0xFE, 0xF9, 0xC3))
        _text(slide, MARGIN + Inches(0.25), body_top + Inches(0.2),
              W - MARGIN * 2 - Inches(0.5), Inches(0.4),
              f"표본이 {data['total']}건으로 적습니다. 아래 순위는 우연히 만들어졌을 수 있으니 "
              "참고용으로만 보세요.", size=12, bold=True,
              color=RGBColor(0x71, 0x3F, 0x12))
        body_top += Inches(1.0)

    if summary:
        _text(slide, MARGIN, body_top, W - MARGIN * 2, Inches(2.2),
              summary, size=14, spacing=8)
    else:
        _text(slide, MARGIN, body_top, W - MARGIN * 2, Inches(0.5),
              "AI 요약이 포함되지 않았습니다. 아래 슬라이드의 집계를 참고하세요.",
              size=12, color=MUTED)


def _slide_severity(prs, data):
    slide = _blank(prs)
    _title(slide, "심각도 분포", f"기간 내 전체 {data['total']}건")

    rows = [r for r in data["by_severity"] if r["count"] > 0]
    if not rows:
        _text(slide, MARGIN, Inches(2.2), W - MARGIN * 2, Inches(0.5),
              "해당 기간에 이벤트가 없습니다.", size=14, color=MUTED)
        return

    cd = CategoryChartData()
    cd.categories = [r["name"] for r in rows]
    cd.add_series("건수", [r["count"] for r in rows])

    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.8),
        Inches(7.2), Inches(4.6), cd)
    chart = gf.chart
    _style_chart(chart)

    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(11)
    plot.data_labels.font.name = BODY_FONT
    plot.data_labels.font.color.rgb = INK
    # 심각도마다 정해진 색을 쓴다. 화면과 같은 색이어야 한다.
    plot.vary_by_categories = True
    for i, r in enumerate(rows):
        point = plot.series[0].points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = SEVERITY.get(r["name"], ACCENT)

    # 오른쪽에 숫자 표. 색을 구분하기 어려운 경우의 대비책이다.
    total = data["total"] or 1
    table_rows = [
        (r["name"], str(r["count"]), f"{round(r['count'] / total * 100, 1)}%")
        for r in data["by_severity"]
    ]
    _table(slide, Inches(8.3), Inches(2.0), Inches(4.3), table_rows,
           ["심각도", "건수", "비중"],
           col_widths=[Inches(1.8), Inches(1.2), Inches(1.3)])


def _slide_sources(prs, data):
    slide = _blank(prs)
    _title(slide, "출처별 조치 필요 건수",
           f"critical + error · 직전 {data['days']}일과 비교")

    rows = data["by_source"][:8]
    if not rows:
        _text(slide, MARGIN, Inches(2.2), W - MARGIN * 2, Inches(0.5),
              "해당 기간에 조치가 필요한 이벤트가 없습니다.", size=14, color=MUTED)
        return

    cd = CategoryChartData()
    cd.categories = [r["name"] for r in rows]
    cd.add_series("직전", [r["prev"] for r in rows])
    cd.add_series("이번", [r["count"] for r in rows])

    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.8),
        Inches(7.2), Inches(4.6), cd)
    chart = gf.chart
    _style_chart(chart, legend=True)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor(0xC5, 0xCC, 0xD4)
    chart.series[1].format.fill.solid()
    chart.series[1].format.fill.fore_color.rgb = ACCENT

    table_rows = []
    for r in rows:
        txt, colour = _delta_text(r["change"])
        table_rows.append((r["name"], str(r["count"]), str(r["prev"]), (txt, colour)))
    _table(slide, Inches(8.3), Inches(2.0), Inches(4.3), table_rows,
           ["출처", "이번", "직전", "변화"],
           col_widths=[Inches(1.6), Inches(0.9), Inches(0.9), Inches(0.9)])


def _slide_noise(prs, data):
    """반복 알람과 조용해진 출처. 둘 다 '들여다볼 후보' 성격이라 한 장에 묶는다."""
    if not data["repeated"] and not data["silent_sources"]:
        return
    slide = _blank(prs)
    _title(slide, "확인이 필요한 것들")

    top = Inches(1.8)
    half = (W - MARGIN * 2 - Inches(0.4)) / 2

    _text(slide, MARGIN, top, half, Inches(0.3),
          "반복 알람 (노이즈 후보)", size=15, bold=True, font=HEAD_FONT)
    if data["repeated"]:
        rows = [(str(r["count"]), r["severity"], r["sample"][:40])
                for r in data["repeated"][:8]]
        _table(slide, MARGIN, top + Inches(0.45), half, rows,
               ["횟수", "심각도", "예시"],
               col_widths=[Inches(0.8), Inches(1.2), half - Inches(2.0)])
    else:
        _text(slide, MARGIN, top + Inches(0.5), half, Inches(0.6),
              "같은 지문이 두 번 이상 나온 알람이 없습니다.",
              size=12, color=MUTED)

    right = MARGIN + half + Inches(0.4)
    _text(slide, right, top, half, Inches(0.3),
          "조용해진 출처", size=15, bold=True, font=HEAD_FONT)
    if data["silent_sources"]:
        _text(slide, right, top + Inches(0.45), half, Inches(0.8),
              "직전 기간엔 이벤트가 있었으나 이번엔 0건입니다.\n"
              "안정된 것일 수도, 수집이 끊긴 것일 수도 있습니다.",
              size=11, color=MUTED, spacing=3)
        _card(slide, right, top + Inches(1.35), half, Inches(2.2))
        _text(slide, right + Inches(0.25), top + Inches(1.55), half - Inches(0.5), Inches(1.9),
              "\n".join(f"· {s}" for s in data["silent_sources"][:10]), size=13, spacing=5)
    else:
        _text(slide, right, top + Inches(0.5), half, Inches(0.6),
              "조용해진 출처가 없습니다.", size=12, color=MUTED)


def _slide_resources(prs, rdiff):
    if not rdiff or not rdiff.get("changes"):
        return
    slide = _blank(prs)
    s = rdiff["summary"]
    _title(slide, "인프라 변경",
           f"스냅샷 #{rdiff['base']['snapshot_id']} → #{rdiff['target']['snapshot_id']} · "
           f"생성 {s['added']} · 변경 {s['modified']} · 삭제 {s['removed']}")

    label = {"added": "생성", "modified": "변경", "removed": "삭제"}
    colour = {"added": ACCENT, "modified": SEVERITY["warning"], "removed": SEVERITY["critical"]}
    rows = []
    for c in rdiff["changes"][:11]:
        detail = ", ".join(f"{f['field']}: {f['before']} → {f['after']}"
                           for f in c["fields"]) or "—"
        rows.append((
            (label[c["change"]], colour[c["change"]]),
            c["resource_id"],
            c["resource_type"],
            detail[:70],
        ))
    _table(slide, MARGIN, Inches(1.85), W - MARGIN * 2, rows,
           ["구분", "리소스", "종류", "변경 내용"],
           col_widths=[Inches(0.9), Inches(2.3), Inches(2.2), Inches(6.5)])


def _slide_caveats(prs, data):
    """마무리. 이 숫자들을 어떻게 읽어야 하는지 적는다."""
    slide = _blank(prs)
    _bg(slide, prs, DARK_BG)
    _text(slide, MARGIN, Inches(0.9), W - MARGIN * 2, Inches(0.7),
          "읽을 때 주의할 점", size=32, bold=True, color=DARK_INK, font=HEAD_FONT)

    notes = [
        ("알람이 몰린 곳이 원인이 아닐 수 있습니다",
         "DB 가 느려지면 그걸 호출하는 웹 서버들이 타임아웃 알람을 쏩니다. "
         "집계는 피해자를 범인으로 지목하는 경향이 있습니다."),
        ("조용한 것과 건강한 것은 다릅니다",
         "알람은 설정된 것만 발생합니다. 모니터링이 걸려 있지 않은 영역은 "
         "아무리 아파도 0건입니다."),
        ("절대 건수보다 변화를 보세요",
         "임계값 정책이 다르면 시스템 간 건수 비교는 의미가 약합니다. "
         "직전 기간 대비 증감이 더 믿을 만한 신호입니다."),
    ]
    if data["small_sample"]:
        notes.insert(0, (
            f"표본이 {data['total']}건으로 적습니다",
            "건수가 적으면 아무 문제 없는 인프라에서도 그럴듯한 1등이 만들어집니다. "
            "이번 리포트의 순위는 참고용으로만 보세요."))

    top = Inches(1.9)
    for title, body in notes:
        _text(slide, MARGIN, top, W - MARGIN * 2, Inches(0.35),
              title, size=17, bold=True, color=ACCENT, font=HEAD_FONT)
        _text(slide, MARGIN, top + Inches(0.42), W - MARGIN * 2 - Inches(1.5), Inches(0.7),
              body, size=13, color=DARK_MUTED)
        top += Inches(1.25)


# ----------------------------------------------------------------------
# 진입점
# ----------------------------------------------------------------------
def build(data, rdiff=None, summary=None):
    """리포트 자료를 받아 pptx 를 만들고 BytesIO 로 돌려준다."""
    prs = Presentation()
    # python-pptx 기본은 4:3 이다. 16:9 로 바꾸지 않으면 좌우가 잘린다.
    prs.slide_width, prs.slide_height = W, H

    _slide_cover(prs, data)
    _slide_summary(prs, data, summary)
    _slide_severity(prs, data)
    _slide_sources(prs, data)
    _slide_noise(prs, data)
    _slide_resources(prs, rdiff)
    _slide_caveats(prs, data)

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
